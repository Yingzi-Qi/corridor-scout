"""Step 02 — validate the schema and select the newest applicable period."""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from .settings import DESTINATIONS, PROVIDERS, REQUIRED_COLUMNS, SHEET_NAME, period_key


def open_validated_sheet(path: Path):
    """Open the expected worksheet and fail clearly if required columns changed."""
    workbook = load_workbook(path, read_only=True, data_only=True)
    if SHEET_NAME not in workbook.sheetnames:
        raise ValueError(f"Required worksheet missing: {SHEET_NAME}")
    sheet = workbook[SHEET_NAME]
    headers = next(sheet.iter_rows(values_only=True))
    column = {value: index for index, value in enumerate(headers) if value}
    missing = sorted(REQUIRED_COLUMNS - set(column))
    if missing:
        workbook.close()
        raise ValueError(f"Required columns missing: {', '.join(missing)}")
    return workbook, sheet, column


def validate_source(path: Path, requested_period: str | None = None) -> dict:
    """Scan the source and return the validated period and row count."""
    workbook, sheet, column = open_validated_sheet(path)
    periods = set()
    source_rows_scanned = 0
    for row in sheet.iter_rows(min_row=2, values_only=True):
        source_rows_scanned += 1
        if (
            row[column["source_code"]] == "GBR"
            and row[column["firm"]] in PROVIDERS
            and row[column["destination_name"]] in DESTINATIONS
        ):
            periods.add(str(row[column["period"]]))
    workbook.close()

    if not periods:
        raise ValueError("No UK records matched the configured providers and destinations")
    selected_period = requested_period or max(periods, key=period_key)
    if selected_period not in periods:
        raise ValueError(f"Requested period is unavailable: {selected_period}")
    return {
        "selectedPeriod": selected_period,
        "sourceRowsScanned": source_rows_scanned,
        "schemaStatus": "passed",
    }
