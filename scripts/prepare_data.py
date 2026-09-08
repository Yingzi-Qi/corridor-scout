#!/usr/bin/env python3
"""Prepare a small, auditable Corridor Scout dataset from the World Bank workbook."""

from __future__ import annotations

import argparse
import json
from pathlib import Path

from openpyxl import load_workbook


PROVIDERS = {"Wise", "Remitly", "WorldRemit", "Western Union", "MoneyGram"}
DESTINATIONS = {
    "Bangladesh",
    "Ghana",
    "India",
    "Kenya",
    "Pakistan",
    "Philippines",
    "South Africa",
    "Tanzania",
    "Thailand",
    "Uganda",
}
SPEED_DAYS = {
    "Less than one hour": 0.04,
    "Same day": 0.5,
    "Next day": 1,
    "2 days": 2,
    "3-5 days": 5,
    "6 days or more": 6,
}


def number(value):
    return round(float(value), 6) if value not in (None, "", "..") else None


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()

    workbook = load_workbook(args.workbook, read_only=True, data_only=True)
    sheet = workbook["Dataset (from Q2 2016)"]
    source_rows = sheet.iter_rows(values_only=True)
    headers = next(source_rows)
    column = {value: index for index, value in enumerate(headers) if value}

    offers = []
    for row in source_rows:
        if not (
            row[column["period"]] == "2025_3Q"
            and row[column["source_code"]] == "GBR"
            and row[column["firm"]] in PROVIDERS
            and row[column["destination_name"]] in DESTINATIONS
            and str(row[column["transparent"]]).lower() == "yes"
        ):
            continue

        for tier in ("cc1", "cc2"):
            local_amount = number(row[column[f"{tier} lcu amount"]])
            fee = number(row[column[f"{tier} lcu fee"]])
            total_cost_pct = number(row[column[f"{tier} total cost %"]])
            fx_margin_pct = number(row[column[f"{tier} fx margin"]])
            if None in (local_amount, fee, total_cost_pct, fx_margin_pct):
                continue

            offers.append(
                {
                    "id": f"{row[column['id']]}-{tier}",
                    "sourceRowId": str(row[column["id"]]),
                    "period": "Q3 2025",
                    "origin": "United Kingdom",
                    "destination": row[column["destination_name"]],
                    "provider": row[column["firm"]],
                    "providerType": row[column["firm_type"]],
                    "fundingMethod": row[column["payment instrument"]],
                    "accessPoint": row[column["access point"]],
                    "receiveMethod": row[column["pickup method"]],
                    "speed": row[column["speed actual"]],
                    "speedDays": SPEED_DAYS.get(row[column["speed actual"]], 99),
                    "benchmarkUsd": number(row[column[f"{tier} denomination amount"]]),
                    "sendAmountGbp": local_amount,
                    "feeGbp": fee,
                    "feePct": round(fee / local_amount * 100, 4),
                    "fxMarginPct": fx_margin_pct,
                    "totalCostPct": total_cost_pct,
                    "estimatedTotalCostGbp": round(local_amount * total_cost_pct / 100, 2),
                    "providerFxRate": number(row[column[f"{tier} lcu fx rate"]]),
                    "interbankFxRate": number(row[column["inter lcu bank fx"]]),
                    "collectionDate": (
                        row[column["date"]].date().isoformat()
                        if hasattr(row[column["date"]], "date")
                        else str(row[column["date"]])
                    ),
                    "coverage": row[column["receiving network coverage"]],
                    "note": row[column["Standard Note"]] or "",
                }
            )

    offers.sort(
        key=lambda item: (
            item["destination"],
            item["benchmarkUsd"],
            item["provider"],
            item["totalCostPct"],
        )
    )
    payload = {
        "metadata": {
            "title": "World Bank Remittance Prices Worldwide",
            "sourceUrl": "https://remittanceprices.worldbank.org/data-download",
            "methodologyUrl": "https://remittanceprices.worldbank.org/methodology",
            "license": "CC BY 4.0",
            "period": "Q3 2025",
            "origin": "United Kingdom",
            "scope": "Five providers across ten common UK-origin corridors",
            "recordCount": len(offers),
            "providers": sorted(PROVIDERS),
            "destinations": sorted(DESTINATIONS),
        },
        "offers": offers,
    }
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(payload, ensure_ascii=False, separators=(",", ":")))
    print(f"Wrote {len(offers)} offers to {args.output}")


if __name__ == "__main__":
    main()
